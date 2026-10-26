"""Excel export of a three-way match result."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FILLS = {
    "PASS": PatternFill("solid", fgColor="DCFCE7"),
    "FAIL": PatternFill("solid", fgColor="FEE2E2"),
    "WARN": PatternFill("solid", fgColor="FEF3C7"),
    "INFO": PatternFill("solid", fgColor="E0F2FE"),
}
HEAD_FILL = PatternFill("solid", fgColor="1E293B")
HEAD_FONT = Font(bold=True, color="FFFFFF")


def _sheet(ws, headers, rows, status_col=None, widths=None):
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    for r in rows:
        ws.append(r)
        if status_col is not None:
            st = r[status_col]
            if st in FILLS:
                ws.cell(row=ws.max_row, column=status_col + 1).fill = FILLS[st]
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)


def build_report(result: dict, run_meta: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Three-Way Match Report"
    ws["A1"].font = Font(bold=True, size=14)
    info = [
        ("Run ID", run_meta.get("id")), ("Run at", run_meta.get("created_at")),
        ("Run by", run_meta.get("username")),
        ("Purchase Order", result.get("po_no")), ("Invoice(s)", result.get("invoice_no")),
        ("GRN(s)", result.get("grn_no")), ("Overall result", result.get("overall")),
        ("Checks passed", result["counts"]["PASS"]), ("Failures", result["counts"]["FAIL"]),
        ("Warnings", result["counts"]["WARN"]),
    ]
    for k, v in result.get("settings", {}).items():
        info.append((k.replace("_", " ").title(), v))
    for i, (k, v) in enumerate(info, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)
    overall_fill = {"MATCHED": "PASS", "MATCHED WITH WARNINGS": "WARN", "INCOMPLETE": "WARN"}.get(result["overall"], "FAIL")
    ws.cell(row=9, column=2).fill = FILLS[overall_fill]
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 26, 40

    _sheet(wb.create_sheet("Header Checks"),
           ["Group", "Check", "PO", "Invoice", "GRN", "Status", "Note"],
           [[c["group"], c["check"], _s(c["po"]), _s(c["invoice"]), _s(c["grn"]), c["status"], c["note"]]
            for c in result["header_checks"]],
           status_col=5, widths=[14, 38, 22, 26, 22, 9, 50])

    _sheet(wb.create_sheet("Line Checks"),
           ["PO Line", "Item", "Description", "PO Qty", "GRN Qty", "Inv Qty", "PO Price", "GRN Price",
            "Inv Price", "PO Total", "GRN Total", "Inv Total", "Status", "Issues", "Quantities by document"],
           [[ln["line"], ln["item"], ln["desc"], ln["po_qty"], ln["grn_qty"], ln["inv_qty"], ln["po_price"],
             ln["grn_price"], ln["inv_price"], ln["po_total"], ln["grn_total"], ln["inv_total"],
             ln["status"], "; ".join(ln["issues"]), "; ".join(ln.get("sources") or [])]
            for ln in result["line_checks"]],
           status_col=12, widths=[8, 12, 20, 9, 9, 9, 9, 9, 9, 11, 11, 11, 9, 50, 40])

    if result.get("pairs"):
        _sheet(wb.create_sheet("GRN-Invoice Pairs"),
               ["GRN", "GRN Date", "Received Value", "Invoice", "Invoice Date", "Invoiced Value", "Status",
                "Pair Result", "Findings"],
               [[_s(p["grn"]), _s(p["grn_date"]), _s(p["grn_value"]), _s(p["invoice"]), _s(p["inv_date"]),
                 _s(p["inv_value"]), p["status"], p["verdict"], "; ".join(p["issues"])] for p in result["pairs"]],
               status_col=6, widths=[12, 12, 14, 20, 12, 14, 9, 22, 70])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _s(v):
    return "" if v is None else v
